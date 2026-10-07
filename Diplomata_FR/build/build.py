#!/usr/bin/env python3
"""원고(.txt) -> 낭독 교재 docx (+ PDF).

원고 형식
  @series  시리즈 표시줄            (예: EL AGREGADO · 외교관 시리즈)
  @lang    언어 표시줄              (예: 스페인어 낭독 교재)
  @part    1
  @ep      1
  @title   El atunero retenido
  @ko      억류된 참치잡이배
  @setting 한 줄 배경
  @level   C1
  @cast    NAME | 설명           (여러 줄)
  @note    머리말 안내문          (여러 줄)
  ## 섹션 제목
  NAME: 외국어 ||| 한국어           (대사 1턴)
  + 표현 = 설명                    (섹션 끝 주석, 여러 줄)
  @next    다음 화 예고
"""
import re
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

FONT = "Noto Sans KR"
NAVY = RGBColor(0x1F, 0x2A, 0x44)
GRAY = RGBColor(0x55, 0x55, 0x55)
ACCENT = RGBColor(0x8A, 0x1C, 0x2B)


def set_font(run, size, bold=False, color=None, italic=False):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade(par, fill):
    ppr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)


def fmt(par, before=0, after=2, line=1.15, align=None, left=None):
    pf = par.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if left is not None:
        pf.left_indent = Cm(left)
    if align:
        par.alignment = align


def parse(path):
    m = {"cast": [], "note": [], "body": []}
    cur = None
    for raw in Path(path).read_text(encoding="utf-8").splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("@"):
            key, _, val = line[1:].partition(" ")
            val = val.strip()
            if key in ("cast", "note"):
                m[key].append(val)
            else:
                m[key] = val
        elif line.startswith("## "):
            cur = {"title": line[3:].strip(), "turns": [], "notes": []}
            m["body"].append(cur)
        elif line.startswith("+ "):
            cur["notes"].append(line[2:].strip())
        else:
            mm = re.match(r"^([^:|]{1,30}):\s*(.+?)\s*\|\|\|\s*(.+)$", line)
            if not mm:
                raise SystemExit(f"형식 오류: {line[:80]}")
            cur["turns"].append((mm.group(1).strip(), mm.group(2).strip(), mm.group(3).strip()))
    return m


def build(src, out_docx):
    m = parse(src)
    n_turns = sum(len(s["turns"]) for s in m["body"])
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(1.9)
    st = d.styles["Normal"]
    st.font.name = FONT
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    def para(text="", size=10.5, bold=False, color=None, italic=False, **kw):
        p = d.add_paragraph()
        fmt(p, **{k: v for k, v in kw.items() if k in ("before", "after", "line", "align", "left")})
        if text:
            set_font(p.add_run(text), size, bold, color, italic)
        return p

    para(m.get("series", ""), 10, True, ACCENT, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(m.get("lang", ""), 10, False, GRAY, after=6, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(f"{m['part']}부 · {m['ep']}화", 13, True, NAVY, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(m["title"], 22, True, NAVY, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(m["ko"], 12, False, GRAY, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(m.get("setting", ""), 10, False, GRAY, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(f"대사 {n_turns}턴 · 난이도 {m.get('level', 'C1')} · 지문 없이 대사만으로 구성 · 본문 + 한국어 대역 + 주석",
         9.5, False, GRAY, after=10, align=WD_ALIGN_PARAGRAPH.CENTER)

    para("등장인물", 12, True, NAVY, before=4, after=3)
    for c in m["cast"]:
        name, _, desc = c.partition("|")
        p = para(after=1, left=0.3)
        set_font(p.add_run(name.strip() + "  "), 10, True)
        set_font(p.add_run(desc.strip()), 10, False, GRAY)
    if m["note"]:
        p = para(before=6, after=8)
        shade(p, "F3EEE6")
        set_font(p.add_run("※ " + " ".join(m["note"])), 9.5, False, GRAY)

    counter = 0
    for s in m["body"]:
        p = para(s["title"], 13, True, NAVY, before=14, after=5)
        pPr = p._p.get_or_add_pPr()
        bd = OxmlElement("w:pBdr")
        bt = OxmlElement("w:bottom")
        for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "8A1C2B")):
            bt.set(qn("w:" + k), v)
        bd.append(bt)
        pPr.append(bd)
        for who, es, ko in s["turns"]:
            counter += 1
            p = para(before=4, after=0, left=0.0)
            p.paragraph_format.keep_with_next = True
            set_font(p.add_run(f"{counter:03d}  "), 7.5, False, GRAY)
            set_font(p.add_run(who + "   "), 10.5, True, ACCENT)
            set_font(p.add_run(es), 11)
            p2 = para(ko, 9.5, False, GRAY, after=2, left=1.05)
        if s["notes"]:
            p = para("주석 · 표현 노트", 9.5, True, NAVY, before=8, after=1)
            shade(p, "EEF1F7")
            for n in s["notes"]:
                p = para(after=0, left=0.3)
                shade(p, "EEF1F7")
                k, _, v = n.partition("=")
                set_font(p.add_run("• " + k.strip()), 9.5, True)
                if v:
                    set_font(p.add_run("  —  " + v.strip()), 9.5, False, GRAY)

    if m.get("next"):
        para("다음 화 예고", 12, True, NAVY, before=16, after=3)
        p = para(m["next"], 10, False, GRAY, after=0)
        shade(p, "F3EEE6")

    # 쪽번호
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run()
    set_font(r, 8, False, GRAY)
    for t, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if t:
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), t)
        else:
            e = OxmlElement("w:instrText")
            e.set(qn("xml:space"), "preserve")
            e.text = txt
        r._r.append(e)

    Path(out_docx).parent.mkdir(parents=True, exist_ok=True)
    d.save(out_docx)
    return n_turns


if __name__ == "__main__":
    src = Path(sys.argv[1])
    out_dir = Path(sys.argv[2])
    stem = sys.argv[3]
    docx = out_dir / "docx" / f"{stem}.docx"
    n = build(src, docx)
    print(f"{stem}: {n}턴")
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir",
                    str(out_dir / "pdf"), str(docx)], check=True,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
